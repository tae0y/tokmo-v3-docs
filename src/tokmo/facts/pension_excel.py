"""Strict pension Excel workbook parser (M1, M2, m3, m6).

Pinned openpyxl 3.1.5 + defusedxml 0.7.1 + et-xmlfile 2.0.0 (M2).
ZIP-entry pass before openpyxl load (M1). Workbook-content pass after load.
Privacy boundary: only typed holdings data leaves this module (§1.8).
"""

import hashlib
import io
import re
import stat
import unicodedata
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from uuid import uuid4
from zoneinfo import ZoneInfo

# M2: parser library pins asserted at import time (if/raise, not bare assert).
import openpyxl

from tokmo.values import canonical_decimal
from tokmo.values.current_holdings import HoldingItemV1
from tokmo.values.pension_holdings import (
    PENSION_HOLDINGS_SCHEMA_VERSION,
    PENSION_PARSER_VERSION,
    PensionHoldingsRevisionV1,
    pension_reconciliation_digest,
    pension_semantic_digest,
)

if openpyxl.__version__ != "3.1.5":
    raise ValueError(f"openpyxl version pin required: got {openpyxl.__version__}")
if not openpyxl.DEFUSEDXML:
    raise ValueError("defusedxml backend required")
import defusedxml

if defusedxml.__version__ != "0.7.1":
    raise ValueError(f"defusedxml version pin required: got {defusedxml.__version__}")

_KST = ZoneInfo("Asia/Seoul")

# Workbook bounds (hard constants §1.2)
WORKBOOK_MAX_BYTES = 10 * 1024 * 1024
WORKBOOK_MAX_ZIP_ENTRIES = 200
WORKBOOK_MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
WORKBOOK_MAX_ROWS = 1000
WORKBOOK_MAX_COLS = 32
WORKBOOK_MAX_CELL_STRING_LENGTH = 256
WORKBOOK_HEADER_ROW = 1
WORKBOOK_DATA_START_ROW = 2

FORBIDDEN_ZIP_ENTRIES = frozenset(
    {
        "xl/vbaProject.bin",
    }
)
FORBIDDEN_ZIP_PREFIXES = frozenset(
    {
        "xl/media/",
        "xl/charts/",
        "xl/drawings/",
        "xl/embeddings/",
        "xl/externalLinks/",
    }
)

HEADER_ALIASES_V1: dict[str, frozenset[str]] = {
    "instrument_name": frozenset(
        {
            "instrument_name",
            "종목명",
            "상품명",
        }
    ),
    "symbol": frozenset(
        {
            "symbol",
            "단축코드",
            "종목코드",
            "표준코드",
        }
    ),
    "market": frozenset(
        {
            "market",
            "시장",
            "거래소",
        }
    ),
    "currency": frozenset(
        {
            "currency",
            "통화",
            "거래통화",
        }
    ),
    "quantity": frozenset(
        {
            "quantity",
            "수량",
            "보유수량",
        }
    ),
    "average_cost": frozenset(
        {
            "average_cost",
            "평균단가",
            "평단가",
            "매입단가",
        }
    ),
    "evaluation_amount": frozenset(
        {
            "evaluation_amount",
            "평가금액",
        }
    ),
    "current_price": frozenset(
        {
            "current_price",
            "현재가",
        }
    ),
    "source_as_of": frozenset(
        {
            "source_as_of",
            "기준일",
            "기준일자",
            "as_of_date",
        }
    ),
}

REQUIRED_FIELDS = frozenset(
    {
        "instrument_name",
        "symbol",
        "market",
        "currency",
        "quantity",
        "average_cost",
        "source_as_of",
    }
)
OPTIONAL_FIELDS = frozenset({"evaluation_amount", "current_price"})

# KR symbols: 6-digit numeric (e.g. 005930) or 6-char alphanumeric with leading
# zero preserved (e.g. 0167A0, 0008T0). Must be string-typed in the cell.
_KR_SYMBOL_RE = re.compile(r"^0[0-9A-Z]{5}$|^[0-9]{6}$")
_US_SYMBOL_RE = re.compile(r"^[A-Z0-9][A-Z0-9.-]{0,19}$")


class PensionExcelParseError(ValueError):
    """Typed rejection from the parser."""

    def __init__(self, reason: str):
        self.reason = reason
        super().__init__(reason)


def _normalize_header(value: str) -> str:
    return unicodedata.normalize("NFC", value).casefold().strip()


def _parse_source_as_of(cell) -> datetime:
    val = cell.value
    if isinstance(val, datetime):
        if val.tzinfo is None or val.utcoffset() is None:
            # Excel stores datetimes naive; interpret as KST local
            return val.replace(tzinfo=_KST)
        return val.astimezone(_KST)
    if isinstance(val, date):
        return datetime(val.year, val.month, val.day, tzinfo=_KST)
    raise PensionExcelParseError("invalid_source_as_of_type")


def _check_landing_file(path: Path) -> bytes:
    """Phase 0: file system checks + read raw bytes.

    M3: check is_symlink() on the original path BEFORE resolve();
    reject st_nlink != 1 (hardlink).
    """
    if path.is_symlink():
        raise PensionExcelParseError("symlink_rejected")
    resolved = path.resolve()
    info = resolved.stat()
    if not stat.S_ISREG(info.st_mode):
        raise PensionExcelParseError("not_regular_file")
    if info.st_nlink != 1:
        raise PensionExcelParseError("hardlink_rejected")
    if stat.S_IMODE(info.st_mode) != 0o600:
        raise PensionExcelParseError("invalid_file_mode")
    if info.st_size > WORKBOOK_MAX_BYTES:
        raise PensionExcelParseError("oversized_file")
    raw = resolved.read_bytes()
    if len(raw) > WORKBOOK_MAX_BYTES:
        raise PensionExcelParseError("oversized_file")
    return raw


def _zip_entry_pass(raw: bytes) -> None:
    """Phase 1: ZIP-entry rejection before openpyxl load (M1)."""
    if not raw.startswith(b"PK\x03\x04"):
        if raw.startswith(b"\xd0\xcf\x11\xe0"):
            raise PensionExcelParseError("encrypted_or_ole_format")
        raise PensionExcelParseError("not_zip_format")
    try:
        zf = zipfile.ZipFile(io.BytesIO(raw))
    except zipfile.BadZipFile:
        raise PensionExcelParseError("corrupt_zip")
    entries = zf.infolist()
    if len(entries) > WORKBOOK_MAX_ZIP_ENTRIES:
        raise PensionExcelParseError("too_many_zip_entries")
    total = sum(e.file_size for e in entries)
    if total > WORKBOOK_MAX_UNCOMPRESSED_BYTES:
        raise PensionExcelParseError("oversized_uncompressed")
    for entry in entries:
        name = entry.filename
        if name in FORBIDDEN_ZIP_ENTRIES:
            if name == "xl/vbaProject.bin":
                raise PensionExcelParseError("macro_detected")
        for prefix in FORBIDDEN_ZIP_PREFIXES:
            if name.startswith(prefix):
                reason = {
                    "xl/media/": "media_detected",
                    "xl/charts/": "chart_detected",
                    "xl/drawings/": "drawing_detected",
                    "xl/embeddings/": "embedding_detected",
                    "xl/externalLinks/": "external_link_detected",
                }[prefix]
                raise PensionExcelParseError(reason)


def _load_workbook_from_bytes(raw: bytes):
    """Phase 2: workbook-content pass after openpyxl load.

    M3: load from BytesIO(raw) so exactly the hashed bytes are parsed,
    eliminating the TOCTOU between the ZIP/hash pass and the openpyxl load.
    """
    import openpyxl.utils.datetime as opxl_dt

    wb = openpyxl.load_workbook(
        io.BytesIO(raw),
        data_only=False,
        read_only=False,
        keep_vba=False,
        keep_links=False,
        rich_text=False,
    )
    if len(wb.worksheets) != 1:
        raise PensionExcelParseError("multi_sheet")
    ws = wb.worksheets[0]
    if ws.sheet_state != "visible":
        raise PensionExcelParseError("hidden_sheet")
    if wb.epoch is not opxl_dt.CALENDAR_WINDOWS_1900:
        raise PensionExcelParseError("date_system_1904")
    return wb, ws


def _validate_workbook_content(ws) -> None:
    """Phase 2 content checks: formulas, hidden rows, merged cells."""
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0
    if max_row > WORKBOOK_MAX_ROWS:
        raise PensionExcelParseError("too_many_rows")
    if max_col > WORKBOOK_MAX_COLS:
        raise PensionExcelParseError("too_many_cols")
    for row in ws.iter_rows(
        min_row=WORKBOOK_DATA_START_ROW,
        max_row=max_row,
        max_col=max_col,
    ):
        for cell in row:
            if cell.data_type == "f":
                raise PensionExcelParseError("formula_detected")
    for r in range(WORKBOOK_DATA_START_ROW, max_row + 1):
        rd = ws.row_dimensions.get(r)
        if rd is not None and rd.hidden:
            raise PensionExcelParseError("hidden_row_detected")
    for merged_range in ws.merged_cells.ranges:
        if merged_range.min_row >= WORKBOOK_DATA_START_ROW:
            raise PensionExcelParseError("merged_cell_detected")


def _parse_headers(ws) -> dict[str, int]:
    """Parse and validate the header row. Returns {canonical_field: col_idx_0based}."""
    max_col = ws.max_column or 0
    raw_headers = []
    for c in range(1, max_col + 1):
        cell = ws.cell(row=WORKBOOK_HEADER_ROW, column=c)
        val = cell.value
        if not isinstance(val, str) or not val.strip():
            raise PensionExcelParseError("empty_header")
        if len(val) > WORKBOOK_MAX_CELL_STRING_LENGTH:
            raise PensionExcelParseError("oversized_cell_string")
        raw_headers.append(_normalize_header(val))

    seen = set()
    field_map: dict[str, int] = {}
    for idx, hdr in enumerate(raw_headers):
        matched = False
        for canonical_field, aliases in HEADER_ALIASES_V1.items():
            if hdr in aliases or hdr == canonical_field:
                if canonical_field in field_map:
                    raise PensionExcelParseError("duplicate_header")
                if hdr in seen:
                    raise PensionExcelParseError("duplicate_header")
                field_map[canonical_field] = idx
                seen.add(hdr)
                matched = True
                break
        if not matched:
            raise PensionExcelParseError("unmapped_header")
    for req in REQUIRED_FIELDS:
        if req not in field_map:
            raise PensionExcelParseError("missing_required_field")
    return field_map


def _validate_symbol(value, market: str) -> str:
    if not isinstance(value, str):
        raise PensionExcelParseError("symbol_must_be_string")
    sym = value.strip()
    if not sym:
        raise PensionExcelParseError("empty_symbol")
    if len(sym) > 64:
        raise PensionExcelParseError("oversized_symbol")
    if market == "KR":
        if not _KR_SYMBOL_RE.fullmatch(sym):
            raise PensionExcelParseError("invalid_kr_symbol")
    elif market == "US":
        if not _US_SYMBOL_RE.fullmatch(sym):
            raise PensionExcelParseError("invalid_us_symbol")
    else:
        raise PensionExcelParseError("invalid_market")
    return sym


def _validate_market_currency(market_val, currency_val) -> tuple[str, str]:
    if market_val not in ("KR", "US"):
        raise PensionExcelParseError("invalid_market")
    expected_cur = "KRW" if market_val == "KR" else "USD"
    if currency_val != expected_cur:
        raise PensionExcelParseError("currency_mismatch")
    return market_val, currency_val


def _parse_cell_decimal(value) -> str:
    try:
        d = Decimal(str(value))
        return canonical_decimal(d)
    except (InvalidOperation, ValueError):
        raise PensionExcelParseError("invalid_decimal")


def parse_pension_excel(
    path: Path,
    *,
    now: datetime,
    explicit_source_as_of: datetime | None = None,
) -> PensionHoldingsRevisionV1:
    """Parse a pension Excel workbook and return a typed candidate revision.

    This function does NOT mutate the file, write to any store, or access
    any network. It reads the file at *path* and returns a typed DTO.

    M4: For a zero-row (header-only / LIQUIDATED) workbook, the caller
    must supply *explicit_source_as_of*. For data rows, all rows must
    carry an identical source_as_of or reject with inconsistent_source_as_of.
    """
    raw = _check_landing_file(path)
    source_content_sha256 = hashlib.sha256(raw).hexdigest()
    _zip_entry_pass(raw)

    wb, ws = _load_workbook_from_bytes(raw)
    try:
        _validate_workbook_content(ws)

        # Sheet name validation: ^pension_holdings_v(\d+)$
        sheet_name = ws.title
        m = re.match(r"^pension_holdings_v(\d+)$", sheet_name)
        if m is None:
            raise PensionExcelParseError("invalid_sheet_name")
        sheet_schema_version = int(m.group(1))
        if sheet_schema_version != 1:
            raise PensionExcelParseError("unsupported_sheet_version")

        field_map = _parse_headers(ws)

        max_row = ws.max_row or 0
        holdings_list = []
        names = []
        eval_amounts = {}
        cur_prices = {}
        seen_keys = set()

        for row_idx in range(WORKBOOK_DATA_START_ROW, max_row + 1):
            market_val = ws.cell(row=row_idx, column=field_map["market"] + 1).value
            currency_val = ws.cell(row=row_idx, column=field_map["currency"] + 1).value
            market, currency = _validate_market_currency(market_val, currency_val)

            symbol_val = ws.cell(row=row_idx, column=field_map["symbol"] + 1).value
            symbol = _validate_symbol(symbol_val, market)

            qty_cell = ws.cell(row=row_idx, column=field_map["quantity"] + 1)
            avg_cell = ws.cell(row=row_idx, column=field_map["average_cost"] + 1)
            quantity = Decimal(_parse_cell_decimal(qty_cell.value))
            avg_cost = Decimal(_parse_cell_decimal(avg_cell.value))

            if quantity <= 0:
                raise PensionExcelParseError("zero_quantity_row")
            if avg_cost < 0:
                raise PensionExcelParseError("negative_cost")

            name_cell = ws.cell(row=row_idx, column=field_map["instrument_name"] + 1)
            name_val = name_cell.value
            if not isinstance(name_val, str) or not name_val.strip():
                raise PensionExcelParseError("invalid_instrument_name")
            if len(name_val) > WORKBOOK_MAX_CELL_STRING_LENGTH:
                raise PensionExcelParseError("oversized_cell_string")

            as_of_cell = ws.cell(row=row_idx, column=field_map["source_as_of"] + 1)
            source_as_of = _parse_source_as_of(as_of_cell)

            dup_key = (market, symbol, currency)
            if dup_key in seen_keys:
                raise PensionExcelParseError("duplicate_symbol")
            seen_keys.add(dup_key)

            # Optional fields
            if "evaluation_amount" in field_map:
                eval_cell = ws.cell(
                    row=row_idx, column=field_map["evaluation_amount"] + 1
                )
                if eval_cell.value is not None:
                    eval_amounts[symbol] = _parse_cell_decimal(eval_cell.value)
            if "current_price" in field_map:
                price_cell = ws.cell(row=row_idx, column=field_map["current_price"] + 1)
                if price_cell.value is not None:
                    cur_prices[symbol] = _parse_cell_decimal(price_cell.value)

            market_raw = market
            market_field = "market"

            holding = HoldingItemV1(
                symbol=symbol,
                quantity=str(quantity),
                avg_price=str(avg_cost),
                currency=currency,
                market_raw=market_raw,
                market_field=market_field,
            )
            holdings_list.append(holding)
            names.append((holding, name_val.strip()))

        holdings = tuple(holdings_list)
        # M4: all data rows must carry identical source_as_of.
        # For empty workbook, require explicit_source_as_of.
        if holdings:
            row_source_as_ofs = []
            for row_idx in range(WORKBOOK_DATA_START_ROW, max_row + 1):
                as_of_cell = ws.cell(row=row_idx, column=field_map["source_as_of"] + 1)
                row_source_as_ofs.append(_parse_source_as_of(as_of_cell))
            first = row_source_as_ofs[0]
            for other in row_source_as_ofs[1:]:
                if other != first:
                    raise PensionExcelParseError("inconsistent_source_as_of")
            source_as_of_final = first
            liquidation_flag = 0
        else:
            # m3: zero-row workbook → LIQUIDATED.
            if explicit_source_as_of is None:
                raise PensionExcelParseError("missing_source_as_of")
            source_as_of_final = explicit_source_as_of.astimezone(_KST)
            liquidation_flag = 1

        semantic = pension_semantic_digest(holdings)
        reconciliation = pension_reconciliation_digest(
            tuple(names),
            eval_amounts or None,
            cur_prices or None,
        )
        revision_id = uuid4().hex
        uploaded_at = now
        validated_at = now

        return PensionHoldingsRevisionV1(
            revision_id=revision_id,
            schema_version=PENSION_HOLDINGS_SCHEMA_VERSION,
            source_content_sha256=source_content_sha256,
            parser_version=PENSION_PARSER_VERSION,
            sheet_schema_version=sheet_schema_version,
            holdings=holdings,
            source_as_of=source_as_of_final,
            uploaded_at=uploaded_at,
            validated_at=validated_at,
            semantic_digest=semantic,
            reconciliation_digest=reconciliation,
            holding_count=len(holdings),
            liquidation_flag=liquidation_flag,
        )
    finally:
        wb.close()
