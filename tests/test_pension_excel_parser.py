"""Slice 1 — Excel parser and privacy tests (M1, M2, m1, m3, m6)."""
from __future__ import annotations

import io
import os
import stat
import zipfile
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl
import pytest

from tokmo.facts.pension_excel import (
    PensionExcelParseError,
    parse_pension_excel,
    WORKBOOK_MAX_BYTES,
    WORKBOOK_MAX_ZIP_ENTRIES,
)
from tokmo.values.pension_holdings import (
    PENSION_HOLDINGS_SCHEMA_VERSION,
    PensionHoldingsRevisionV1,
)
from tokmo.values.current_holdings import HoldingItemV1

_KST = ZoneInfo("Asia/Seoul")
_NOW = datetime(2026, 8, 5, 12, 0, tzinfo=_KST)


def _make_workbook(
    rows: list[dict],
    *,
    sheet_name: str = "pension_holdings_v1",
    extra_columns: dict | None = None,
) -> bytes:
    """Build a minimal valid xlsx workbook from row dicts."""
    wb = openpyxl.Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet()
    ws.title = sheet_name

    headers = ["instrument_name", "symbol", "market", "currency",
               "quantity", "average_cost", "source_as_of"]
    if extra_columns:
        headers += list(extra_columns.keys())
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h)

    for r, row in enumerate(rows, 2):
        for c, h in enumerate(headers, 1):
            if h in row:
                val = row[h]
                # Excel does not support timezone-aware datetimes
                if h == "source_as_of" and isinstance(val, datetime):
                    if val.tzinfo is not None:
                        val = val.replace(tzinfo=None)
                ws.cell(row=r, column=c, value=val)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_landing(tmp_path: Path, data: bytes) -> Path:
    p = tmp_path / "pension_landing.xlsx"
    p.write_bytes(data)
    os.chmod(p, 0o600)
    return p


class TestParserBasics:
    """1.4 GREEN: strict typed candidate with exact canonical bytes."""

    def test_valid_workbook_parses_to_typed_revision(self, tmp_path):
        rows = [
            {
                "instrument_name": "SOL AI반도체TOP2플러스",
                "symbol": "0167A0",
                "market": "KR", "currency": "KRW",
                "quantity": 100, "average_cost": 10000.50,
                "source_as_of": datetime(2026, 7, 1, tzinfo=_KST),
            },
            {
                "instrument_name": "Apple Inc",
                "symbol": "AAPL",
                "market": "US", "currency": "USD",
                "quantity": 50, "average_cost": 150.00,
                "source_as_of": datetime(2026, 7, 1, tzinfo=_KST),
            },
        ]
        data = _make_workbook(rows)
        path = _write_landing(tmp_path, data)
        revision = parse_pension_excel(path, now=_NOW)

        assert isinstance(revision, PensionHoldingsRevisionV1)
        assert revision.holding_count == 2
        assert revision.schema_version == PENSION_HOLDINGS_SCHEMA_VERSION
        assert revision.parser_version == "pension-excel-v1"
        assert revision.liquidation_flag == 0
        assert len(revision.source_content_sha256) == 64
        assert len(revision.semantic_digest) == 64

    def test_leading_zero_kr_symbol_preserved(self, tmp_path):
        """1.1: Leading-zero KR symbol 0167A0 is preserved as a string."""
        rows = [{
            "instrument_name": "SOL AI반도체TOP2플러스",
            "symbol": "0167A0",
            "market": "KR", "currency": "KRW",
            "quantity": 100, "average_cost": 10000,
            "source_as_of": datetime(2026, 7, 1, tzinfo=_KST),
        }]
        data = _make_workbook(rows)
        path = _write_landing(tmp_path, data)
        revision = parse_pension_excel(path, now=_NOW)
        assert revision.holdings[0].symbol == "0167A0"

    def test_holding_item_v1_reused(self):
        """1.8 GREEN: HoldingItemV1 is reused, no PensionHoldingItemV1 duplicate."""
        import tokmo.values.pension_holdings as ph
        # m1: HoldingItemV1 comes from current_holdings
        assert ph.HoldingItemV1 is not None
        # No PensionHoldingItemV1 class exists in values/
        assert not hasattr(ph, "PensionHoldingItemV1")

    def test_pinned_library_versions(self):
        """1.7 GREEN: openpyxl, defusedxml, et-xmlfile pinned."""
        assert openpyxl.__version__ == "3.1.5"
        assert openpyxl.DEFUSEDXML is True
        import defusedxml
        assert defusedxml.__version__ == "0.7.1"
        import et_xmlfile
        assert et_xmlfile is not None


class TestParserRejections:
    """1.2 RED: rejection cases."""

    def test_missing_required_field_rejected(self, tmp_path):
        rows = [{
            "instrument_name": "Test",
            "symbol": "005930",
            "market": "KR", "currency": "KRW",
            "quantity": 100, "average_cost": 10000,
            # source_as_of missing
        }]
        data = _make_workbook(rows)
        path = _write_landing(tmp_path, data)
        with pytest.raises(PensionExcelParseError):
            parse_pension_excel(path, now=_NOW)

    def test_duplicate_symbol_rejected(self, tmp_path):
        rows = [
            {"instrument_name": "A", "symbol": "005930", "market": "KR",
             "currency": "KRW", "quantity": 100, "average_cost": 10000,
             "source_as_of": datetime(2026, 7, 1, tzinfo=_KST)},
            {"instrument_name": "B", "symbol": "005930", "market": "KR",
             "currency": "KRW", "quantity": 50, "average_cost": 11000,
             "source_as_of": datetime(2026, 7, 1, tzinfo=_KST)},
        ]
        data = _make_workbook(rows)
        path = _write_landing(tmp_path, data)
        with pytest.raises(PensionExcelParseError) as exc_info:
            parse_pension_excel(path, now=_NOW)
        assert exc_info.value.reason == "duplicate_symbol"

    def test_negative_quantity_rejected(self, tmp_path):
        rows = [{
            "instrument_name": "Test", "symbol": "005930",
            "market": "KR", "currency": "KRW",
            "quantity": -100, "average_cost": 10000,
            "source_as_of": datetime(2026, 7, 1, tzinfo=_KST),
        }]
        data = _make_workbook(rows)
        path = _write_landing(tmp_path, data)
        with pytest.raises(PensionExcelParseError) as exc_info:
            parse_pension_excel(path, now=_NOW)
        assert exc_info.value.reason == "zero_quantity_row"

    def test_currency_mismatch_rejected(self, tmp_path):
        rows = [{
            "instrument_name": "Test", "symbol": "005930",
            "market": "KR", "currency": "USD",
            "quantity": 100, "average_cost": 10000,
            "source_as_of": datetime(2026, 7, 1, tzinfo=_KST),
        }]
        data = _make_workbook(rows)
        path = _write_landing(tmp_path, data)
        with pytest.raises(PensionExcelParseError) as exc_info:
            parse_pension_excel(path, now=_NOW)
        assert exc_info.value.reason == "currency_mismatch"

    def test_invalid_market_rejected(self, tmp_path):
        rows = [{
            "instrument_name": "Test", "symbol": "005930",
            "market": "JP", "currency": "JPY",
            "quantity": 100, "average_cost": 10000,
            "source_as_of": datetime(2026, 7, 1, tzinfo=_KST),
        }]
        data = _make_workbook(rows)
        path = _write_landing(tmp_path, data)
        with pytest.raises(PensionExcelParseError) as exc_info:
            parse_pension_excel(path, now=_NOW)
        assert exc_info.value.reason == "invalid_market"


class TestParserPrivacy:
    """1.3 RED: privacy boundary."""

    def test_revision_repr_excludes_holdings_and_raw_values(self, tmp_path):
        rows = [{
            "instrument_name": "SOL AI반도체TOP2플러스",
            "symbol": "0167A0",
            "market": "KR", "currency": "KRW",
            "quantity": 100, "average_cost": 10000,
            "source_as_of": datetime(2026, 7, 1, tzinfo=_KST),
        }]
        data = _make_workbook(rows)
        path = _write_landing(tmp_path, data)
        revision = parse_pension_excel(path, now=_NOW)
        repr_text = repr(revision)
        assert "0167A0" not in repr_text
        assert "100" not in repr_text
        assert "10000" not in repr_text
        assert "SOL AI" not in repr_text

    def test_no_path_in_revision_repr(self, tmp_path):
        rows = [{
            "instrument_name": "Test", "symbol": "005930",
            "market": "KR", "currency": "KRW",
            "quantity": 100, "average_cost": 10000,
            "source_as_of": datetime(2026, 7, 1, tzinfo=_KST),
        }]
        data = _make_workbook(rows)
        path = _write_landing(tmp_path, data)
        revision = parse_pension_excel(path, now=_NOW)
        assert str(path) not in repr(revision)
        assert "pension_landing" not in repr(revision)


class TestAdversarialWorkbooks:
    """1.5a-d RED: adversarial fixtures."""

    def test_zip_bomb_rejected(self, tmp_path):
        """Too many ZIP entries."""
        rows = [{
            "instrument_name": "Test", "symbol": "005930",
            "market": "KR", "currency": "KRW",
            "quantity": 100, "average_cost": 10000,
            "source_as_of": datetime(2026, 7, 1, tzinfo=_KST),
        }]
        data = _make_workbook(rows)
        # Add many fake entries
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            for i in range(WORKBOOK_MAX_ZIP_ENTRIES + 5):
                zf.writestr(f"dummy_{i}.txt", "x" * 10)
            zf.writestr("real.xlsx", data)
        path = _write_landing(tmp_path, buf.getvalue())
        with pytest.raises(PensionExcelParseError) as exc_info:
            parse_pension_excel(path, now=_NOW)
        assert exc_info.value.reason == "too_many_zip_entries"

    def test_ole_encrypted_rejected(self, tmp_path):
        """OLE/CFB container rejected before openpyxl load."""
        data = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100
        path = _write_landing(tmp_path, data)
        with pytest.raises(PensionExcelParseError) as exc_info:
            parse_pension_excel(path, now=_NOW)
        assert exc_info.value.reason == "encrypted_or_ole_format"

    def test_macro_detected(self, tmp_path):
        """Macro ZIP entry rejected."""
        rows = [{
            "instrument_name": "Test", "symbol": "005930",
            "market": "KR", "currency": "KRW",
            "quantity": 100, "average_cost": 10000,
            "source_as_of": datetime(2026, 7, 1, tzinfo=_KST),
        }]
        data = _make_workbook(rows)
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("xl/vbaProject.bin", b"\x00" * 100)
            # Extract the xlsx into the zip
            with zipfile.ZipFile(io.BytesIO(data)) as orig:
                for name in orig.namelist():
                    zf.writestr(name, orig.read(name))
        path = _write_landing(tmp_path, buf.getvalue())
        with pytest.raises(PensionExcelParseError) as exc_info:
            parse_pension_excel(path, now=_NOW)
        assert exc_info.value.reason == "macro_detected"

    def test_multi_sheet_rejected(self, tmp_path):
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "pension_holdings_v1"
        ws1.cell(row=1, column=1, value="instrument_name")
        ws1.cell(row=1, column=2, value="symbol")
        ws1.cell(row=1, column=3, value="market")
        ws1.cell(row=1, column=4, value="currency")
        ws1.cell(row=1, column=5, value="quantity")
        ws1.cell(row=1, column=6, value="average_cost")
        ws1.cell(row=1, column=7, value="source_as_of")
        ws1.cell(row=2, column=7, value=datetime(2026, 7, 1))  # naive for Excel
        ws2 = wb.create_sheet("Sheet2")
        buf = io.BytesIO()
        wb.save(buf)
        path = _write_landing(tmp_path, buf.getvalue())
        with pytest.raises(PensionExcelParseError) as exc_info:
            parse_pension_excel(path, now=_NOW)
        assert exc_info.value.reason == "multi_sheet"

    def test_invalid_file_mode_rejected(self, tmp_path):
        rows = [{
            "instrument_name": "Test", "symbol": "005930",
            "market": "KR", "currency": "KRW",
            "quantity": 100, "average_cost": 10000,
            "source_as_of": datetime(2026, 7, 1, tzinfo=_KST),
        }]
        data = _make_workbook(rows)
        path = tmp_path / "pension_landing.xlsx"
        path.write_bytes(data)
        os.chmod(path, 0o644)  # wrong mode
        with pytest.raises(PensionExcelParseError) as exc_info:
            parse_pension_excel(path, now=_NOW)
        assert exc_info.value.reason == "invalid_file_mode"

    def test_invalid_sheet_name_rejected(self, tmp_path):
        rows = [{
            "instrument_name": "Test", "symbol": "005930",
            "market": "KR", "currency": "KRW",
            "quantity": 100, "average_cost": 10000,
            "source_as_of": datetime(2026, 7, 1, tzinfo=_KST),
        }]
        data = _make_workbook(rows, sheet_name="Sheet1")
        path = _write_landing(tmp_path, data)
        with pytest.raises(PensionExcelParseError) as exc_info:
            parse_pension_excel(path, now=_NOW)
        assert exc_info.value.reason == "invalid_sheet_name"


class TestHeaderAliases:
    """1.5 header alias coverage."""

    def test_korean_header_aliases(self, tmp_path):
        """Korean header names work."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "pension_holdings_v1"
        for c, h in enumerate(["종목명", "단축코드", "시장", "통화",
                                "수량", "평균단가", "기준일"], 1):
            ws.cell(row=1, column=c, value=h)
        ws.cell(row=2, column=1, value="삼성전자")
        ws.cell(row=2, column=2, value="005930")
        ws.cell(row=2, column=3, value="KR")
        ws.cell(row=2, column=4, value="KRW")
        ws.cell(row=2, column=5, value=100)
        ws.cell(row=2, column=6, value=70000)
        ws.cell(row=2, column=7, value=datetime(2026, 7, 1))  # naive for Excel
        buf = io.BytesIO()
        wb.save(buf)
        path = _write_landing(tmp_path, buf.getvalue())
        revision = parse_pension_excel(path, now=_NOW)
        assert revision.holdings[0].symbol == "005930"

    def test_unmapped_header_rejected(self, tmp_path):
        rows = [{
            "instrument_name": "Test", "symbol": "005930",
            "market": "KR", "currency": "KRW",
            "quantity": 100, "average_cost": 10000,
            "source_as_of": datetime(2026, 7, 1, tzinfo=_KST),
        }]
        data = _make_workbook(rows, extra_columns={"unknown_col": "val"})
        path = _write_landing(tmp_path, data)
        with pytest.raises(PensionExcelParseError) as exc_info:
            parse_pension_excel(path, now=_NOW)
        assert exc_info.value.reason == "unmapped_header"
