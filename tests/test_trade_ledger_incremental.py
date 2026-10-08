"""Synthetic period exports: batch identity, not invented cumulative row IDs."""
import hashlib
import json
import sqlite3
from dataclasses import replace
from datetime import date, datetime, timezone
from uuid import uuid4

import pytest

from tokmo.records.trade_ledger_canonical import canonical_payload_bytes
from tokmo.records.trade_ledger_store import TradeLedgerStore
from tokmo.values.trade_ledger import ImportAuditMeta, TradeLedgerEvent

SOURCE = 'retirement_excel_delta_v1'
SCOPE = 'retirement_pension_1'


def event(day='2026-09-14', price='1000'):
    return TradeLedgerEvent(
        source=SOURCE, account_scope=SCOPE, source_record_key='batch',
        event_kind='EXECUTION', supersedes_source_record_key=None,
        trade_date=date.fromisoformat(day), transaction_type_raw='ETF매수',
        instrument_raw='SYNTH-ETF', quantity='2', unit_price=price,
        amount=str(int(price) * 2), medium_raw='SYNTH', currency_raw='KRW')


def batch(*events):
    ordered = sorted(events, key=canonical_payload_bytes)
    digest = hashlib.sha256(b'pension-delta-v1\0' + b'\n'.join(
        canonical_payload_bytes(e) for e in ordered)).hexdigest()
    return tuple(replace(e, source_record_key=f'batch-{digest}-row-{i:06d}')
                 for i, e in enumerate(ordered, 1))


def store(path):
    return TradeLedgerStore(path, clock=lambda: datetime.now(timezone.utc),
                            id_factory=lambda: uuid4().hex)


def append(s, events):
    return s.append_batch(events, import_audit=ImportAuditMeta('ab'*32, 'delta-test'))


def test_disjoint_delta_and_replay_preserve_existing(tmp_path):
    s = store(tmp_path / 'ledger.sqlite3')
    old = replace(event('2026-07-01'), source='retirement_excel_v1', source_record_key='row-000001')
    assert append(s, (old,)).inserted_count == 1
    delta = batch(event(), event('2026-09-17'))
    result = append(s, delta)
    assert result.status == 'SAVED' and result.inserted_count == 2
    replay = append(s, tuple(reversed(delta)))
    assert replay.status == 'SAVED' and replay.duplicate_count == 2
    assert replay.inserted_count == 0


@pytest.mark.parametrize('days', [('2026-09-15',), ('2026-09-10','2026-09-15'), ('2026-09-17','2026-09-21')])
def test_overlapping_nonidentical_batch_refused_atomically(tmp_path, days):
    s = store(tmp_path / 'ledger.sqlite3')
    assert append(s, batch(event(), event('2026-09-17'))).inserted_count == 2
    result = append(s, batch(*(event(d, '1200') for d in days)))
    assert result.status == 'FAILED' and result.inserted_count == 0
    assert result.reason == 'batch_identity_conflict'


def test_equal_fills_preserved_and_changed_identity_refused(tmp_path):
    s = store(tmp_path / 'ledger.sqlite3')
    delta = batch(event(), event())
    assert append(s, delta).inserted_count == 2
    assert append(s, batch(event(), event())).duplicate_count == 2
    corrupt = (replace(delta[0], unit_price='9'), delta[1])
    assert append(s, corrupt).status == 'FAILED'


def test_legacy_contract_remains_full_history(tmp_path):
    s = store(tmp_path / 'ledger.sqlite3')
    old = replace(event(), source='retirement_excel_v1', source_record_key='row-000001')
    assert append(s, (old,)).status == 'SAVED'
    assert append(s, (replace(old, source_record_key='row-000002'),)).status == 'FAILED'


def test_legacy_resubmission_cannot_double_count_delta(tmp_path):
    s = store(tmp_path / 'ledger.sqlite3')
    old = replace(event('2026-09-13'), source='retirement_excel_v1', source_record_key='row-000001')
    assert append(s, (old,)).inserted_count == 1
    assert append(s, batch(event())).inserted_count == 1
    newer = replace(event(), source='retirement_excel_v1', source_record_key='row-000002')
    assert append(s, (old, newer)).status == 'FAILED'
    with sqlite3.connect(tmp_path/'ledger.sqlite3') as c:
        assert c.execute('select count(*) from trade_ledger_events').fetchone()[0] == 2


def test_invalid_member_and_concurrent_replay_are_atomic(tmp_path):
    from concurrent.futures import ThreadPoolExecutor
    s = store(tmp_path / 'ledger.sqlite3'); rows = batch(event(), event('2026-09-17'))
    assert append(s, (rows[0], replace(rows[1], quantity=None))).status == 'FAILED'
    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(lambda _: append(s, rows), range(2)))
    assert sum(r.inserted_count for r in results) == 2
    assert sorted(r.duplicate_count for r in results) == [0, 2]


@pytest.mark.parametrize('alteration', ['same', 'repack', 'reorder'])
def test_xlsx_entrypoint_append_then_semantic_replay(tmp_path, alteration):
    import subprocess
    import sys
    from pathlib import Path
    from openpyxl import Workbook
    from tests.test_pension_operational_tools import batch_environment
    registry, snapshots, _, ledger_path = batch_environment(tmp_path)
    before = registry.read_active_revision()
    def workbook(path, rows):
        wb = Workbook(); ws = wb.active; ws.title = 'Sheet1'
        ws.append(['일자','거래구분','종목명','수량','단가','금액','거래매체'])
        for row in rows: ws.append(row)
        wb.save(path); wb.close(); path.chmod(0o600)
    rows = [['2026/09/14','ETF매수','SYNTH-ETF',2,1000,2000,'SYNTH'],
            ['2026/09/17','ETF매도','SYNTH-ETF',1,1200,1200,'SYNTH']]
    file = tmp_path / 'delta.xlsx'; workbook(file, rows)
    tool = Path(__file__).resolve().parents[1] / 'tools/import_pension_delta_excel.py'
    def run(path):
        return subprocess.run([sys.executable, str(tool), '--file', str(path),
            '--root', str(tmp_path), '--apply'], capture_output=True, text=True)
    first = run(file)
    assert first.returncode == 0, first.stderr
    assert json.loads(first.stdout)['appended_count'] == 2
    n = len(snapshots.list_pension_capture_headers())
    if alteration != 'same':
        file = tmp_path / 'renamed.xlsx'
        workbook(file, rows[::-1] if alteration == 'reorder' else rows)
    second = run(file)
    assert second.returncode == 0, second.stderr
    assert json.loads(second.stdout)['outcome'] == 'DUPLICATE_NOOP'
    assert len(snapshots.list_pension_capture_headers()) == n
    assert registry.read_active_revision() == before
    with sqlite3.connect(ledger_path) as c:
        assert c.execute('select count(*) from trade_ledger_events').fetchone()[0] == 2


@pytest.mark.parametrize('bad', ['formula', 'header'])
def test_xlsx_rejects_untrusted_shape_without_store(tmp_path, bad):
    import subprocess
    import sys
    from pathlib import Path
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    ws.append(['일자','거래구분','종목명','수량','단가','금액','거래매체'])
    ws.append(['2026/09/14','ETF매수','SYNTH',1,100,100,'SYNTH'])
    ws['D2' if bad == 'formula' else 'A1'] = '=1+1' if bad == 'formula' else 'WRONG'
    path = tmp_path/'bad.xlsx'; wb.save(path); wb.close(); path.chmod(0o600)
    root = tmp_path/'absent'
    tool = Path(__file__).resolve().parents[1]/'tools/import_pension_delta_excel.py'
    r = subprocess.run([sys.executable,str(tool),'--file',str(path),'--root',str(root),'--apply'],capture_output=True,text=True)
    assert r.returncode != 0
    assert json.loads(r.stdout)['outcome'] == 'REFUSED'
    assert not root.exists()
