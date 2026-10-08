"""Account-status workbook parser + reviewed classification contract.

Requirement 1: the parser produces a V1 candidate holding exactly the verified
ETF rows and excluding the typed non-ETF rows.
Requirement 2: an unknown/unclassified row or any count mismatch fails closed —
there is no heuristic bond-as-ETF admission anywhere.
Requirement 8: workbook hardening, Decimal provenance, parser privacy.

Every fixture is synthetic. No real instrument name, quantity, price, amount,
account identifier, hash, path, or workbook metadata appears in this file.
"""
from __future__ import annotations

import hashlib
import io
import os
import zipfile
from dataclasses import replace
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl
import pytest

from tokmo.facts import pension_account_status_excel as target
from tokmo.facts import pension_excel as v1_parser
from tokmo.facts.pension_account_status_excel import (
    ACCOUNT_STATUS_FIELDS,
    ACCOUNT_STATUS_OPTIONAL_FIELDS,
    ACCOUNT_STATUS_REQUIRED_FIELDS,
    ACCOUNT_STATUS_SHEET_NAME,
    PERIOD_SUMMARY_COLUMN_COUNT,
    PERIOD_SUMMARY_FIELDS,
    PensionExcelParseError,
    parse_pension_account_status,
    parse_pension_period_summary,
)
from tokmo.values.current_holdings import HoldingItemV1
from tokmo.values.pension_account_status import (
    CLASSIFICATION_PLAN_VERSION,
    COVERAGE_ETF_ONLY_PARTIAL,
    COVERAGE_ETF_ONLY_WHOLE,
    DISPOSITION_EXCLUDED_NON_ETF,
    DISPOSITION_VERIFIED_KRX_ETF,
    EXCLUSION_CASH_LIKE,
    EXCLUSION_CORPORATE_BOND,
    EXCLUSION_MUTUAL_FUND,
    PENSION_ACCOUNT_STATUS_SHEET_SCHEMA_VERSION,
    PensionClassificationError,
    PensionClassificationPlanV1,
    PensionPeriodSummaryRowV1,
    PensionRowDispositionV1,
    PensionSourceRowV1,
    derive_average_cost,
    execution_field_offenders,
)
from tokmo.values.pension_holdings import PENSION_PARSER_VERSION

_KST = ZoneInfo("Asia/Seoul")
_NOW = datetime(2026, 8, 13, 21, 0, tzinfo=_KST)
_AS_OF = datetime(2026, 8, 13, tzinfo=_KST)


# ---------------------------------------------------------------------------
# Synthetic fixtures — 14 source rows: 10 ETFs + 2 bonds + 1 fund + 1 cash
# ---------------------------------------------------------------------------

def status_rows() -> list[list]:
    rows = [
        [f"SYNTH-ETF-{n:02d}", 10, 1000, 1100, 100, 10, 7]
        for n in range(1, 11)
    ]
    rows.append(["SYNTH-BOND-ALPHA", 2, 200, 210, 10, 5, 1])
    rows.append(["SYNTH-BOND-BETA", 2, 200, 210, 10, 5, 1])
    rows.append(["SYNTH-FUND-BALANCED", 4, 400, 420, 20, 5, 2])
    rows.append(["SYNTH-CASH-RESERVE", 500, 500, 506, 0, 0, 3])
    return rows


def etf_symbol(n: int) -> str:
    return f"90000{n}" if n < 10 else "900010"


def reviewed_plan(
    *,
    etf_ordinals: tuple[int, ...] = tuple(range(1, 11)),
    excluded: tuple[tuple[int, str], ...] = (
        (11, EXCLUSION_CORPORATE_BOND), (12, EXCLUSION_CORPORATE_BOND),
        (13, EXCLUSION_MUTUAL_FUND), (14, EXCLUSION_CASH_LIKE),
    ),
) -> PensionClassificationPlanV1:
    dispositions = []
    for ordinal in etf_ordinals:
        dispositions.append(PensionRowDispositionV1(
            ordinal=ordinal,
            disposition=DISPOSITION_VERIFIED_KRX_ETF,
            symbol=etf_symbol(ordinal),
            exclusion_kind=None,
        ))
    for ordinal, kind in excluded:
        dispositions.append(PensionRowDispositionV1(
            ordinal=ordinal,
            disposition=DISPOSITION_EXCLUDED_NON_ETF,
            symbol=None,
            exclusion_kind=kind,
        ))
    dispositions.sort(key=lambda item: item.ordinal)
    return PensionClassificationPlanV1(
        plan_version=CLASSIFICATION_PLAN_VERSION,
        expected_row_count=len(dispositions),
        expected_etf_count=len(etf_ordinals),
        expected_excluded_count=len(excluded),
        rows=tuple(dispositions),
    )


def build_status_workbook(
    rows: list[list] | None = None,
    *,
    headers: list[str] | None = None,
    sheet_name: str = ACCOUNT_STATUS_SHEET_NAME,
) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for column, header in enumerate(
            headers if headers is not None else list(ACCOUNT_STATUS_FIELDS), 1):
        sheet.cell(row=1, column=column, value=header)
    for index, row in enumerate(status_rows() if rows is None else rows, 2):
        for column, value in enumerate(row, 1):
            sheet.cell(row=index, column=column, value=value)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def summary_rows() -> list[list]:
    return [
        [f"SYNTH-ETF-{n:02d}", 10, 4, 100, 110, 1000, 440, 2, 38, 9.5]
        for n in range(1, 6)
    ]


def build_summary_workbook(
    rows: list[list] | None = None,
    *,
    leaf_headers: list[str] | None = None,
    merge_group_header: bool = False,
    merge_data_cell: bool = False,
) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = ACCOUNT_STATUS_SHEET_NAME
    sheet.cell(row=1, column=1, value="SYNTH-GROUP")
    for column, header in enumerate(
            leaf_headers if leaf_headers is not None
            else list(PERIOD_SUMMARY_FIELDS), 1):
        sheet.cell(row=2, column=column, value=header)
    for index, row in enumerate(summary_rows() if rows is None else rows, 3):
        for column, value in enumerate(row, 1):
            sheet.cell(row=index, column=column, value=value)
    if merge_group_header:
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
    if merge_data_cell:
        sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=2)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def landing(tmp_path: Path, data: bytes, name: str = "landing.xlsx") -> Path:
    path = tmp_path / name
    path.write_bytes(data)
    os.chmod(path, 0o600)
    return path


def parse_status(tmp_path, data=None, plan=None, **kwargs):
    return parse_pension_account_status(
        landing(tmp_path, build_status_workbook() if data is None else data),
        now=_NOW, source_as_of=_AS_OF,
        plan=reviewed_plan() if plan is None else plan, **kwargs,
    )


# ---------------------------------------------------------------------------
# Requirement 1 — exactly the verified ETF rows become holdings
# ---------------------------------------------------------------------------

def test_candidate_holds_exactly_the_ten_verified_etfs(tmp_path):
    result = parse_status(tmp_path)
    assert result.revision.holding_count == 10
    assert len(result.revision.holdings) == 10
    assert result.included_etf_count == 10
    assert result.source_row_count == 14
    assert result.excluded_row_count == 4


def test_excluded_rows_never_appear_as_holdings(tmp_path):
    result = parse_status(tmp_path)
    symbols = {item.symbol for item in result.revision.holdings}
    assert symbols == {etf_symbol(n) for n in range(1, 11)}
    # Nothing bond/fund/cash-shaped leaked in.
    assert len(symbols) == 10


def test_excluded_kind_counts_are_typed_and_truthful(tmp_path):
    result = parse_status(tmp_path)
    assert dict(result.excluded_kind_counts) == {
        EXCLUSION_CORPORATE_BOND: 2,
        EXCLUSION_MUTUAL_FUND: 1,
        EXCLUSION_CASH_LIKE: 1,
    }


def test_coverage_statement_never_claims_a_whole_account(tmp_path):
    result = parse_status(tmp_path)
    assert result.coverage_statement == COVERAGE_ETF_ONLY_PARTIAL
    assert "FULL" not in result.coverage_statement
    assert result.revision.holding_count < result.source_row_count


def test_coverage_is_whole_only_when_nothing_was_excluded(tmp_path):
    plan = reviewed_plan(etf_ordinals=tuple(range(1, 11)), excluded=())
    result = parse_status(
        tmp_path, build_status_workbook(status_rows()[:10]), plan)
    assert result.coverage_statement == COVERAGE_ETF_ONLY_WHOLE
    assert result.excluded_row_count == 0


def test_candidate_is_an_ordinary_v1_revision(tmp_path):
    result = parse_status(tmp_path)
    revision = result.revision
    assert revision.parser_version == PENSION_PARSER_VERSION == "pension-excel-v1"
    assert revision.schema_version == 1
    assert revision.liquidation_flag == 0
    assert revision.source_as_of == _AS_OF
    assert all(isinstance(item, HoldingItemV1) for item in revision.holdings)


def test_admitted_holdings_are_krw_kr_with_derived_price(tmp_path):
    for item in parse_status(tmp_path).revision.holdings:
        assert item.currency == "KRW"
        assert item.market_raw == "KR"
        assert item.market_field == "classification_plan"
        assert item.quantity == "10"
        assert item.avg_price == "100"


def test_zero_etf_plan_is_not_a_liquidation(tmp_path):
    """An ETF-only revision with no ETFs must not claim the account is empty:
    the account may still hold excluded non-ETF assets."""
    plan = reviewed_plan(
        etf_ordinals=(),
        excluded=tuple((n, EXCLUSION_CORPORATE_BOND) for n in range(1, 15)),
    )
    result = parse_status(tmp_path, plan=plan)
    assert result.revision.holding_count == 0
    assert result.revision.liquidation_flag == 0
    assert result.excluded_row_count == 14


def test_source_content_hash_is_the_hash_of_the_parsed_bytes(tmp_path):
    data = build_status_workbook()
    result = parse_status(tmp_path, data)
    assert result.revision.source_content_sha256 == \
        hashlib.sha256(data).hexdigest()


def test_input_shape_is_recorded_without_changing_the_output_contract(tmp_path):
    revision = parse_status(tmp_path).revision
    assert revision.sheet_schema_version == \
        PENSION_ACCOUNT_STATUS_SHEET_SCHEMA_VERSION
    assert revision.sheet_schema_version != 1
    assert revision.parser_version == PENSION_PARSER_VERSION


# ---------------------------------------------------------------------------
# Requirement 2 — no heuristic admission; every mismatch fails closed
# ---------------------------------------------------------------------------

def test_an_unclassified_row_fails_closed(tmp_path):
    """A workbook row the reviewer never saw must block the candidate."""
    plan = reviewed_plan()
    short = replace(
        plan,
        rows=plan.rows[:13],
        expected_row_count=13,
        expected_excluded_count=3,
    )
    with pytest.raises(PensionExcelParseError) as error:
        parse_status(tmp_path, plan=short)
    assert error.value.reason == "row_count_contract_mismatch"


def test_extra_workbook_row_beyond_the_reviewed_plan_fails_closed(tmp_path):
    rows = status_rows() + [["SYNTH-UNREVIEWED-01", 1, 100, 100, 0, 0, 1]]
    with pytest.raises(PensionExcelParseError) as error:
        parse_status(tmp_path, build_status_workbook(rows))
    assert error.value.reason == "row_count_contract_mismatch"


def test_a_bond_is_never_admitted_because_its_row_looks_tradable(tmp_path):
    """The bond rows carry perfectly tradable-looking quantities and amounts.
    Only the reviewed disposition decides, so they stay excluded."""
    result = parse_status(tmp_path)
    assert result.revision.holding_count == 10
    assert result.excluded_row_count == 4

    # And the parser contains no name/code-shape heuristic at all: the only
    # thing that admits a row is the reviewed disposition.
    source = Path(target.__file__).read_text(encoding="utf-8")
    for heuristic in ("startswith(", "endswith(", ".upper()", "re.search",
                      "fullmatch"):
        assert heuristic not in source, heuristic
    assert source.count("DISPOSITION_VERIFIED_KRX_ETF") >= 1


def test_declared_etf_count_mismatch_fails_closed():
    """A plan whose declared counts contradict its own rows cannot exist, so a
    mis-declared count can never reach the parser."""
    rows = reviewed_plan().rows
    with pytest.raises(PensionClassificationError) as error:
        PensionClassificationPlanV1(
            plan_version=CLASSIFICATION_PLAN_VERSION,
            expected_row_count=14,
            expected_etf_count=9,      # rows actually declare 10
            expected_excluded_count=5,
            rows=rows,
        )
    assert error.value.reason == "plan_etf_count_mismatch"


def test_declared_excluded_count_mismatch_fails_closed():
    rows = reviewed_plan().rows
    with pytest.raises(PensionClassificationError) as error:
        PensionClassificationPlanV1(
            plan_version=CLASSIFICATION_PLAN_VERSION,
            expected_row_count=14,
            expected_etf_count=10,
            expected_excluded_count=3,  # rows actually declare 4
            rows=rows,
        )
    assert error.value.reason == "plan_excluded_count_mismatch"


def test_plan_with_duplicate_symbols_is_refused():
    plan = reviewed_plan()
    rows = list(plan.rows)
    rows[1] = replace(rows[1], symbol=rows[0].symbol)
    with pytest.raises(PensionClassificationError) as error:
        PensionClassificationPlanV1(
            plan_version=CLASSIFICATION_PLAN_VERSION,
            expected_row_count=14, expected_etf_count=10,
            expected_excluded_count=4, rows=tuple(rows),
        )
    assert error.value.reason == "plan_duplicate_symbol"


def test_plan_with_ordinal_gaps_is_refused():
    plan = reviewed_plan()
    rows = list(plan.rows)
    rows[13] = replace(rows[13], ordinal=99)
    with pytest.raises(PensionClassificationError) as error:
        PensionClassificationPlanV1(
            plan_version=CLASSIFICATION_PLAN_VERSION,
            expected_row_count=14, expected_etf_count=10,
            expected_excluded_count=4, rows=tuple(rows),
        )
    assert error.value.reason == "plan_ordinals_not_contiguous"


def test_verified_etf_disposition_requires_a_valid_short_code():
    with pytest.raises(ValueError):
        PensionRowDispositionV1(
            ordinal=1, disposition=DISPOSITION_VERIFIED_KRX_ETF,
            symbol=None, exclusion_kind=None)
    for bad in ("12345", "1234567", "SYNTH1", "", "abcdef"):
        with pytest.raises(ValueError):
            PensionRowDispositionV1(
                ordinal=1, disposition=DISPOSITION_VERIFIED_KRX_ETF,
                symbol=bad, exclusion_kind=None)


def test_excluded_disposition_cannot_smuggle_a_symbol():
    with pytest.raises(ValueError):
        PensionRowDispositionV1(
            ordinal=1, disposition=DISPOSITION_EXCLUDED_NON_ETF,
            symbol="900001", exclusion_kind=EXCLUSION_CORPORATE_BOND)


def test_excluded_disposition_requires_a_typed_kind():
    with pytest.raises(ValueError):
        PensionRowDispositionV1(
            ordinal=1, disposition=DISPOSITION_EXCLUDED_NON_ETF,
            symbol=None, exclusion_kind=None)
    with pytest.raises(ValueError):
        PensionRowDispositionV1(
            ordinal=1, disposition=DISPOSITION_EXCLUDED_NON_ETF,
            symbol=None, exclusion_kind="SYNTH_UNKNOWN_KIND")


def test_unknown_disposition_is_refused():
    with pytest.raises(ValueError):
        PensionRowDispositionV1(
            ordinal=1, disposition="ADMIT_ANYWAY",
            symbol="900001", exclusion_kind=None)


def test_a_verified_etf_row_with_non_positive_quantity_fails_closed(tmp_path):
    rows = status_rows()
    rows[0][1] = 0
    with pytest.raises(PensionExcelParseError) as error:
        parse_status(tmp_path, build_status_workbook(rows))
    assert error.value.reason == "non_positive_etf_quantity"


def test_plan_version_is_pinned():
    plan = reviewed_plan()
    with pytest.raises(PensionClassificationError) as error:
        PensionClassificationPlanV1(
            plan_version="pension-account-status-plan-vNEXT",
            expected_row_count=14, expected_etf_count=10,
            expected_excluded_count=4, rows=plan.rows,
        )
    assert error.value.reason == "invalid_plan_version"


def test_a_non_plan_object_is_refused(tmp_path):
    with pytest.raises(PensionExcelParseError) as error:
        parse_pension_account_status(
            landing(tmp_path, build_status_workbook()),
            now=_NOW, source_as_of=_AS_OF, plan={"1": "VERIFIED_KRX_ETF"})
    assert error.value.reason == "invalid_classification_plan"


# ---------------------------------------------------------------------------
# Requirement 8 — header contract, hardening, Decimal provenance, privacy
# ---------------------------------------------------------------------------

def test_header_contract_names_three_required_columns(tmp_path):
    """User-approved contract: 종목명 / 보유량 / 매입금액 are required; the
    evaluation, profit, return and weight columns are optional."""
    assert ACCOUNT_STATUS_REQUIRED_FIELDS == (
        "instrument_repr", "quantity", "purchase_amount")
    assert ACCOUNT_STATUS_OPTIONAL_FIELDS == (
        "evaluation_amount", "profit_loss", "return_rate", "weight")
    assert ACCOUNT_STATUS_FIELDS == (
        ACCOUNT_STATUS_REQUIRED_FIELDS + ACCOUNT_STATUS_OPTIONAL_FIELDS)


def test_korean_header_aliases_are_accepted(tmp_path):
    korean = ["상품명", "수량", "매입금액", "평가금액", "평가손익", "수익률", "비중"]
    result = parse_status(tmp_path, build_status_workbook(headers=korean))
    assert result.revision.holding_count == 10


def test_broker_header_with_boyuryang_and_unknown_weight_is_accepted(tmp_path):
    """The real export header shape (synthetic values): 보유량 is a quantity
    alias and the unrecognised 운용비율 column is ignored, not blocking."""
    broker = ["종목명", "보유량", "매입금액", "평가금액", "평가손익", "수익률", "운용비율"]
    result = parse_status(tmp_path, build_status_workbook(headers=broker))
    assert result.revision.holding_count == 10
    assert result.source_row_count == 14
    for item in result.revision.holdings:
        assert item.quantity == "10"
        assert item.avg_price == "100"


def test_three_required_columns_alone_are_sufficient(tmp_path):
    headers = ["종목명", "보유량", "매입금액"]
    rows = [row[:3] for row in status_rows()]
    result = parse_status(tmp_path, build_status_workbook(rows, headers=headers))
    assert result.revision.holding_count == 10
    assert result.excluded_row_count == 4
    for item in result.revision.holdings:
        assert item.avg_price == derive_average_cost("1000", "10")


def test_dropping_the_trailing_weight_column_is_not_blocking(tmp_path):
    headers = list(ACCOUNT_STATUS_FIELDS)[:6]
    rows = [row[:6] for row in status_rows()]
    result = parse_status(tmp_path, build_status_workbook(rows, headers=headers))
    assert result.revision.holding_count == 10


def test_required_columns_are_mapped_by_name_not_position(tmp_path):
    """Column order carries no meaning; the header name does."""
    headers = ["매입금액", "종목명", "평가금액", "보유량"]
    rows = [[row[2], row[0], row[3], row[1]] for row in status_rows()]
    result = parse_status(tmp_path, build_status_workbook(rows, headers=headers))
    assert result.revision.holding_count == 10
    for item in result.revision.holdings:
        assert item.quantity == "10"
        assert item.avg_price == "100"


def test_reordering_columns_does_not_change_the_digests(tmp_path):
    baseline = parse_status(tmp_path).revision
    headers = list(ACCOUNT_STATUS_FIELDS)
    headers.reverse()
    rows = [list(reversed(row)) for row in status_rows()]
    reordered = parse_status(
        tmp_path, build_status_workbook(rows, headers=headers)).revision
    assert reordered.semantic_digest == baseline.semantic_digest
    assert reordered.reconciliation_digest == baseline.reconciliation_digest


def test_missing_required_column_fails_closed(tmp_path):
    headers = ["종목명", "보유량", "평가금액", "평가손익", "수익률", "운용비율"]
    rows = [[row[0], row[1], row[3], row[4], row[5], row[6]]
            for row in status_rows()]
    with pytest.raises(PensionExcelParseError) as error:
        parse_status(tmp_path, build_status_workbook(rows, headers=headers))
    assert error.value.reason == "missing_required_header"


def test_unknown_header_in_a_required_slot_is_a_missing_required_column(tmp_path):
    unknown = list(ACCOUNT_STATUS_FIELDS)
    unknown[2] = "synthetic_unknown_header"       # purchase_amount slot
    with pytest.raises(PensionExcelParseError) as error:
        parse_status(tmp_path, build_status_workbook(headers=unknown))
    assert error.value.reason == "missing_required_header"


def test_duplicate_required_column_fails_closed(tmp_path):
    headers = ["종목명", "보유량", "매입금액", "수량"]
    rows = [[row[0], row[1], row[2], row[1]] for row in status_rows()]
    with pytest.raises(PensionExcelParseError) as error:
        parse_status(tmp_path, build_status_workbook(rows, headers=headers))
    assert error.value.reason == "duplicate_header"


def test_duplicate_optional_column_fails_closed(tmp_path):
    """A recognised optional column that appears twice is ambiguous, and the
    parser never picks one silently."""
    headers = ["종목명", "보유량", "매입금액", "평가금액", "평가금액"]
    rows = [[row[0], row[1], row[2], row[3], row[3]] for row in status_rows()]
    with pytest.raises(PensionExcelParseError) as error:
        parse_status(tmp_path, build_status_workbook(rows, headers=headers))
    assert error.value.reason == "duplicate_header"


def test_unknown_optional_header_is_ignored(tmp_path):
    unknown = list(ACCOUNT_STATUS_FIELDS)
    unknown[3] = "synthetic_unknown_header"       # evaluation_amount slot
    result = parse_status(tmp_path, build_status_workbook(headers=unknown))
    assert result.revision.holding_count == 10


def test_empty_header_cell_is_still_rejected(tmp_path):
    headers = list(ACCOUNT_STATUS_FIELDS)
    headers[6] = "   "
    with pytest.raises(PensionExcelParseError) as error:
        parse_status(tmp_path, build_status_workbook(headers=headers))
    assert error.value.reason == "empty_header"


def test_blank_optional_cell_is_read_as_absent_not_invented(tmp_path):
    rows = status_rows()
    rows[0][3] = None                              # evaluation_amount blank
    result = parse_status(tmp_path, build_status_workbook(rows))
    assert result.revision.holding_count == 10
    filled = parse_status(tmp_path).revision.reconciliation_digest
    assert result.revision.reconciliation_digest != filled


def test_garbage_in_a_mapped_optional_cell_still_fails_closed(tmp_path):
    rows = status_rows()
    rows[0][3] = "synthetic-not-a-number"
    with pytest.raises(PensionExcelParseError) as error:
        parse_status(tmp_path, build_status_workbook(rows))
    assert error.value.reason == "invalid_decimal"


def test_source_row_optional_fields_default_to_none():
    row = PensionSourceRowV1(
        ordinal=1, instrument_repr="SYNTH-ROW", quantity="10",
        purchase_amount="1000")
    assert row.evaluation_amount is None
    assert row.profit_loss is None
    assert row.return_rate is None
    assert row.weight is None
    with pytest.raises(ValueError):
        PensionSourceRowV1(
            ordinal=1, instrument_repr="SYNTH-ROW", quantity="10",
            purchase_amount="1000", weight="not-a-decimal")


def test_wrong_sheet_name_is_rejected(tmp_path):
    with pytest.raises(PensionExcelParseError) as error:
        parse_status(tmp_path, build_status_workbook(sheet_name="Sheet2"))
    assert error.value.reason == "wrong_sheet_name"


def test_duplicate_instrument_repr_is_rejected(tmp_path):
    rows = status_rows()
    rows[1][0] = rows[0][0]
    with pytest.raises(PensionExcelParseError) as error:
        parse_status(tmp_path, build_status_workbook(rows))
    assert error.value.reason == "duplicate_instrument_repr"


def test_formula_cell_is_rejected(tmp_path):
    rows = status_rows()
    rows[0][3] = "=A1+1"
    with pytest.raises(PensionExcelParseError) as error:
        parse_status(tmp_path, build_status_workbook(rows))
    assert error.value.reason == "formula_detected"


def test_hidden_data_row_is_rejected(tmp_path):
    workbook = openpyxl.load_workbook(io.BytesIO(build_status_workbook()))
    workbook.active.row_dimensions[3].hidden = True
    buffer = io.BytesIO()
    workbook.save(buffer)
    with pytest.raises(PensionExcelParseError) as error:
        parse_status(tmp_path, buffer.getvalue())
    assert error.value.reason == "hidden_row_detected"


def test_merged_data_cell_is_rejected(tmp_path):
    workbook = openpyxl.load_workbook(io.BytesIO(build_status_workbook()))
    workbook.active.merge_cells(
        start_row=2, start_column=1, end_row=2, end_column=2)
    buffer = io.BytesIO()
    workbook.save(buffer)
    with pytest.raises(PensionExcelParseError) as error:
        parse_status(tmp_path, buffer.getvalue())
    assert error.value.reason == "merged_cell_detected"


def test_merge_spilling_from_header_into_data_is_rejected(tmp_path):
    workbook = openpyxl.load_workbook(io.BytesIO(build_status_workbook()))
    workbook.active.merge_cells(
        start_row=1, start_column=1, end_row=2, end_column=1)
    buffer = io.BytesIO()
    workbook.save(buffer)
    with pytest.raises(PensionExcelParseError) as error:
        parse_status(tmp_path, buffer.getvalue())
    assert error.value.reason == "merged_cell_detected"


def test_second_sheet_is_rejected(tmp_path):
    workbook = openpyxl.load_workbook(io.BytesIO(build_status_workbook()))
    workbook.create_sheet("SYNTH-EXTRA")
    buffer = io.BytesIO()
    workbook.save(buffer)
    with pytest.raises(PensionExcelParseError) as error:
        parse_status(tmp_path, buffer.getvalue())
    assert error.value.reason == "multi_sheet"


def test_macro_entry_is_rejected(tmp_path):
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(build_status_workbook())) as source:
        with zipfile.ZipFile(buffer, "w") as sink:
            for item in source.infolist():
                sink.writestr(item, source.read(item.filename))
            sink.writestr("xl/vbaProject.bin", b"synthetic")
    with pytest.raises(PensionExcelParseError) as error:
        parse_status(tmp_path, buffer.getvalue())
    assert error.value.reason == "macro_detected"


def test_external_link_entry_is_rejected(tmp_path):
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(build_status_workbook())) as source:
        with zipfile.ZipFile(buffer, "w") as sink:
            for item in source.infolist():
                sink.writestr(item, source.read(item.filename))
            sink.writestr("xl/externalLinks/externalLink1.xml", b"<x/>")
    with pytest.raises(PensionExcelParseError) as error:
        parse_status(tmp_path, buffer.getvalue())
    assert error.value.reason == "external_link_detected"


def test_encrypted_container_is_rejected(tmp_path):
    with pytest.raises(PensionExcelParseError) as error:
        parse_status(tmp_path, b"\xd0\xcf\x11\xe0synthetic")
    assert error.value.reason == "encrypted_or_ole_format"


def test_non_zip_payload_is_rejected(tmp_path):
    with pytest.raises(PensionExcelParseError) as error:
        parse_status(tmp_path, b"synthetic-not-a-zip")
    assert error.value.reason == "not_zip_format"


def test_world_readable_landing_file_is_rejected(tmp_path):
    path = tmp_path / "loose.xlsx"
    path.write_bytes(build_status_workbook())
    os.chmod(path, 0o644)
    with pytest.raises(PensionExcelParseError) as error:
        parse_pension_account_status(
            path, now=_NOW, source_as_of=_AS_OF, plan=reviewed_plan())
    assert error.value.reason == "invalid_file_mode"


def test_symlinked_landing_file_is_rejected(tmp_path):
    real = landing(tmp_path, build_status_workbook(), "real.xlsx")
    link = tmp_path / "link.xlsx"
    link.symlink_to(real)
    with pytest.raises(PensionExcelParseError) as error:
        parse_pension_account_status(
            link, now=_NOW, source_as_of=_AS_OF, plan=reviewed_plan())
    assert error.value.reason == "symlink_rejected"


def test_hardening_helpers_are_shared_with_the_v1_parser():
    assert target.check_landing_file is v1_parser._check_landing_file
    assert target.zip_entry_pass is v1_parser._zip_entry_pass
    assert target.load_workbook_from_bytes is v1_parser._load_workbook_from_bytes
    assert target.PensionExcelParseError is v1_parser.PensionExcelParseError


# --- Decimal provenance ----------------------------------------------------

def test_average_cost_is_derived_exactly_from_purchase_over_quantity():
    assert derive_average_cost("1000", "10") == "100"
    assert derive_average_cost("1", "3") == "0.3333333333"
    assert derive_average_cost("0", "0") is None


def test_float_cell_does_not_leak_a_binary_expansion(tmp_path):
    rows = status_rows()
    rows[0][1] = 0.5
    rows[0][2] = 0.1
    result = parse_status(tmp_path, build_status_workbook(rows))
    item = result.revision.holdings[0]
    assert item.quantity == "0.5"
    assert item.avg_price == "0.2"
    assert "0000000000000000" not in item.avg_price
    assert Decimal(item.avg_price) == Decimal("0.2")


def test_trailing_zero_amounts_are_canonicalized(tmp_path):
    rows = status_rows()
    rows[0][2] = Decimal("1000.000")
    result = parse_status(tmp_path, build_status_workbook(rows))
    assert result.revision.holdings[0].avg_price == "100"


def test_reconciliation_digest_covers_the_excluded_rows(tmp_path):
    """The excluded assets survive as tamper-evident private evidence, which is
    what makes the user-approved exclusion auditable without a second store."""
    baseline = parse_status(tmp_path).revision.reconciliation_digest
    rows = status_rows()
    rows[10][3] = 999          # change an EXCLUDED bond row's evaluation
    changed = parse_status(
        tmp_path, build_status_workbook(rows)).revision.reconciliation_digest
    assert changed != baseline


def test_semantic_digest_ignores_the_excluded_rows(tmp_path):
    """The public-safe semantic digest is over admitted holdings only."""
    baseline = parse_status(tmp_path).revision.semantic_digest
    rows = status_rows()
    rows[10][3] = 999
    changed = parse_status(
        tmp_path, build_status_workbook(rows)).revision.semantic_digest
    assert changed == baseline


# --- privacy ---------------------------------------------------------------

def test_private_name_is_absent_from_result_and_revision_reprs(tmp_path):
    result = parse_status(tmp_path)
    for rendered in (repr(result), repr(result.revision)):
        for name in ("SYNTH-ETF-01", "SYNTH-BOND-ALPHA",
                     "SYNTH-FUND-BALANCED", "SYNTH-CASH-RESERVE"):
            assert name not in rendered


def test_parse_failure_reasons_carry_no_private_name(tmp_path):
    rows = status_rows()
    rows[1][0] = rows[0][0]
    with pytest.raises(PensionExcelParseError) as error:
        parse_status(tmp_path, build_status_workbook(rows))
    assert "SYNTH" not in error.value.reason


# ---------------------------------------------------------------------------
# Period summary — evidence only
# ---------------------------------------------------------------------------

def test_period_summary_parses_as_aggregate_evidence(tmp_path):
    evidence = parse_pension_period_summary(
        landing(tmp_path, build_summary_workbook()), now=_NOW)
    assert evidence.row_count == 5
    assert evidence.aggregate_only is True
    assert PERIOD_SUMMARY_COLUMN_COUNT == 10


def test_period_summary_carries_no_execution_field(tmp_path):
    evidence = parse_pension_period_summary(
        landing(tmp_path, build_summary_workbook()), now=_NOW)
    assert execution_field_offenders(type(evidence)) == ()
    assert execution_field_offenders(PensionPeriodSummaryRowV1) == ()
    for row in evidence.rows:
        assert not hasattr(row, "trade_date")
        assert not hasattr(row, "source_record_key")
        assert not hasattr(row, "event_kind")


def test_period_summary_accepts_a_merged_group_header(tmp_path):
    evidence = parse_pension_period_summary(
        landing(tmp_path, build_summary_workbook(merge_group_header=True)),
        now=_NOW)
    assert evidence.row_count == 5


def test_period_summary_rejects_a_merged_data_cell(tmp_path):
    with pytest.raises(PensionExcelParseError) as error:
        parse_pension_period_summary(
            landing(tmp_path, build_summary_workbook(merge_data_cell=True)),
            now=_NOW)
    assert error.value.reason == "merged_cell_detected"


def test_period_summary_rejects_an_unknown_leaf_header(tmp_path):
    headers = list(PERIOD_SUMMARY_FIELDS)
    headers[4] = "synthetic_unknown_header"
    with pytest.raises(PensionExcelParseError) as error:
        parse_pension_period_summary(
            landing(tmp_path, build_summary_workbook(leaf_headers=headers)),
            now=_NOW)
    assert error.value.reason == "unmapped_header"


def test_period_summary_is_not_a_holdings_type(tmp_path):
    from tokmo.values.pension_holdings import PensionHoldingsRevisionV1
    evidence = parse_pension_period_summary(
        landing(tmp_path, build_summary_workbook()), now=_NOW)
    assert not isinstance(evidence, PensionHoldingsRevisionV1)
    assert not hasattr(evidence, "holdings")
    assert not hasattr(evidence, "holding_count")
